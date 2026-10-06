from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from fastapi import HTTPException, UploadFile
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from python_api.config import get_settings
from python_api.context import get_request_context
from python_api.embeddings import create_embeddings
from python_api.models import AplusDocument

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".ppt", ".pptx", ".txt", ".md"}


@dataclass(slots=True)
class RawDocument:
    """Describes a source file before text extraction."""

    absolute_path: Path
    source_path: str
    file_name: str
    file_type: str
    document_type: str
    file_size: int
    file_hash: str


class DocumentLoader:
    """Extracts supported office and text formats into metadata-bearing segments."""

    def raw_root(self, tenant: str) -> Path:
        """Returns the configured tenant raw directory."""
        return (get_settings().workspace_root / get_settings().raw_dir / tenant).resolve()

    def scan(self, tenant: str) -> list[RawDocument]:
        """Recursively describes supported tenant files."""
        root = self.raw_root(tenant)
        if not root.exists():
            return []
        return [
            self.describe(path, root, tenant)
            for path in root.rglob("*")
            if path.is_file()
            and path.name not in {".DS_Store", ".WeDrive"}
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
        ]

    def describe(
        self, path: Path, root: Path, tenant: str, file_name: str | None = None
    ) -> RawDocument:
        """Collects stable path, type, size and SHA-256 metadata."""
        name = file_name or path.name
        extension = Path(name).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise HTTPException(400, f"Unsupported file type: {extension or 'none'}")
        source_path = path.relative_to(root).as_posix()
        parts = Path(source_path).parts
        return RawDocument(
            absolute_path=path,
            source_path=source_path,
            file_name=name,
            file_type=extension.removeprefix("."),
            document_type=parts[0] if len(parts) > 1 else "",
            file_size=path.stat().st_size,
            file_hash=hashlib.sha256(path.read_bytes()).hexdigest(),
        )

    def load(self, document: RawDocument) -> list[dict[str, Any]]:
        """Loads one supported file and returns text segments."""
        suffix = document.absolute_path.suffix.lower()
        if suffix in {".txt", ".md"}:
            return [{"text": document.absolute_path.read_text(encoding="utf-8"), "metadata": {}}]
        if suffix == ".pdf":
            from pypdf import PdfReader

            return [
                {
                    "text": self._normalize_pdf(page.extract_text() or ""),
                    "metadata": {"page": index + 1},
                }
                for index, page in enumerate(PdfReader(document.absolute_path).pages)
            ]
        if suffix == ".docx":
            from docx import Document

            loaded = Document(str(document.absolute_path))
            text_content = "\n".join(paragraph.text for paragraph in loaded.paragraphs)
            return [{"text": text_content, "metadata": {}}]
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(document.absolute_path, data_only=True, read_only=True)
            segments = []
            for sheet in workbook.worksheets:
                rows = []
                for row_index, values in enumerate(sheet.iter_rows(values_only=True), 1):
                    cells = [
                        str(value).strip()
                        for value in values
                        if value is not None and str(value).strip()
                    ]
                    if cells:
                        rows.append(f"第 {row_index} 行：" + " | ".join(cells))
                segments.append({"text": "\n".join(rows), "metadata": {"sheet": sheet.title}})
            return segments
        if suffix == ".pptx":
            from pptx import Presentation

            presentation = Presentation(str(document.absolute_path))
            return [
                {
                    "text": "\n".join(
                        shape.text for shape in slide.shapes if hasattr(shape, "text")
                    ),
                    "metadata": {"page": index + 1},
                }
                for index, slide in enumerate(presentation.slides)
            ]
        if suffix == ".ppt":
            content = document.absolute_path.read_bytes().decode("latin1", errors="ignore")
            extracted = re.sub(r"[^\x20-\x7e\u4e00-\u9fa5]+", " ", content)
            extracted = re.sub(r"\s+", " ", extracted).strip()
            if not extracted:
                raise HTTPException(
                    400,
                    "Legacy .ppt file has no extractable text. Convert it to .pptx or PDF and retry.",
                )
            return [{"text": extracted, "metadata": {"legacyPpt": True}}]
        raise HTTPException(400, f"Unsupported file type: {suffix}")

    @staticmethod
    def _normalize_pdf(value: str) -> str:
        value = value.replace("\r\n", "\n")
        value = re.sub(
            r"\b(?:[A-Za-z][ \t]+)+[A-Za-z]\b",
            lambda match: re.sub(r"[ \t]+", "", match.group()),
            value,
        )
        value = re.sub(r"[ \t]+", " ", value)
        return re.sub(r"\n{3,}", "\n\n", value).strip()


class DocumentIngestService:
    """Indexes RAG documents into the Prisma-owned pgvector schema."""

    def __init__(self, db: AsyncSession) -> None:
        self.db = db
        self.settings = get_settings()
        self.loader = DocumentLoader()
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.settings.rag_chunk_size,
            chunk_overlap=self.settings.rag_chunk_overlap,
            separators=["\n\n", "\n", "。", "！", "？", ".", "!", "?", " ", ""],
        )
        self.embeddings = create_embeddings(self.settings)

    async def ingest_rag(self, force: bool = False) -> dict[str, Any]:
        """Scans and indexes the current tenant raw directory."""
        tenant = get_request_context().tenant_code
        return await self._ingest_files(self.loader.scan(tenant), tenant, force)

    async def upload(self, files: list[UploadFile]) -> dict[str, Any]:
        """Validates, stores, deduplicates and indexes a multipart batch."""
        if not files:
            raise HTTPException(400, "files are required")
        if len(files) > self.settings.python_api_max_upload_files:
            raise HTTPException(400, "Too many files")
        tenant = get_request_context().tenant_code
        root = self.loader.raw_root(tenant)
        root.mkdir(parents=True, exist_ok=True)
        hashes: dict[str, str] = {}
        results: list[dict[str, Any]] = []
        for upload in files:
            file_name = Path(upload.filename or "").name.strip()
            if not file_name:
                results.append(
                    {
                        "fileName": upload.filename or "-",
                        "status": "failed",
                        "reason": "invalid file name",
                    }
                )
                continue
            if Path(file_name).suffix.lower() not in SUPPORTED_EXTENSIONS:
                results.append(
                    {
                        "fileName": file_name,
                        "status": "failed",
                        "reason": f"unsupported file type; supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}",
                    }
                )
                continue
            data = await upload.read(self.settings.python_api_max_upload_size_bytes + 1)
            if len(data) > self.settings.python_api_max_upload_size_bytes:
                raise HTTPException(400, "File too large")
            digest = hashlib.sha256(data).hexdigest()
            if digest in hashes:
                results.append(
                    {
                        "fileName": file_name,
                        "status": "duplicate",
                        "reason": f"duplicate of {hashes[digest]} in current batch",
                    }
                )
                continue
            hashes[digest] = file_name
            existing = await self.db.scalar(
                select(AplusDocument).where(
                    AplusDocument.tenant_code == tenant, AplusDocument.file_hash == digest
                )
            )
            if existing:
                results.append(
                    {
                        "fileName": file_name,
                        "status": "duplicate",
                        "reason": "duplicate hash",
                        "duplicateOf": {
                            "documentId": str(existing.id),
                            "fileName": existing.file_name,
                        },
                    }
                )
                continue
            target = self._unique_target(root, file_name)
            target.write_bytes(data)
            try:
                raw = self.loader.describe(target, root, tenant, file_name)
                summary = await self._ingest_files([raw], tenant, False)
                item = summary["files"][0]
                if item.get("status") == "indexed":
                    results.append(
                        {
                            "fileName": file_name,
                            "status": "indexed",
                            "documentId": item["documentId"],
                            "sourcePath": raw.source_path,
                            "chunks": item["chunks"],
                        }
                    )
                else:
                    results.append(
                        {
                            "fileName": file_name,
                            "status": "failed",
                            "sourcePath": raw.source_path,
                            "reason": item.get("reason", "document indexing failed"),
                        }
                    )
            except Exception as error:
                results.append(
                    {
                        "fileName": file_name,
                        "status": "failed",
                        "sourcePath": target.name,
                        "reason": str(error),
                    }
                )
        return {
            "totalFiles": len(files),
            "indexedFiles": sum(item["status"] == "indexed" for item in results),
            "duplicateFiles": sum(item["status"] == "duplicate" for item in results),
            "failedFiles": sum(item["status"] == "failed" for item in results),
            "chunks": sum(item.get("chunks", 0) for item in results),
            "files": results,
        }

    async def _ingest_files(
        self, files: list[RawDocument], tenant: str, force: bool
    ) -> dict[str, Any]:
        summary: dict[str, Any] = {
            "scannedFiles": len(files),
            "indexedFiles": 0,
            "skippedFiles": 0,
            "failedFiles": 0,
            "chunks": 0,
            "files": [],
        }
        for file in files:
            existing = await self.db.scalar(
                select(AplusDocument).where(
                    AplusDocument.tenant_code == tenant,
                    AplusDocument.source_path == file.source_path,
                )
            )
            duplicate = await self.db.scalar(
                select(AplusDocument).where(
                    AplusDocument.tenant_code == tenant, AplusDocument.file_hash == file.file_hash
                )
            )
            if duplicate and (not existing or duplicate.id != existing.id):
                summary["skippedFiles"] += 1
                summary["files"].append(
                    {
                        "sourcePath": file.source_path,
                        "status": "skipped",
                        "reason": "duplicate hash",
                        "duplicateOf": {
                            "documentId": str(duplicate.id),
                            "fileName": duplicate.file_name,
                        },
                    }
                )
                continue
            if (
                existing
                and existing.file_hash == file.file_hash
                and existing.status == "ready"
                and not force
            ):
                summary["skippedFiles"] += 1
                summary["files"].append(
                    {
                        "sourcePath": file.source_path,
                        "status": "skipped",
                        "reason": "hash unchanged",
                    }
                )
                continue
            try:
                chunks = self._build_chunks(file, tenant)
                if not chunks:
                    raise ValueError("no extractable text")
                vectors = await self._embed([item["content"] for item in chunks])
                document_id = await self._replace(file, tenant, chunks, vectors)
                summary["indexedFiles"] += 1
                summary["chunks"] += len(chunks)
                summary["files"].append(
                    {
                        "sourcePath": file.source_path,
                        "status": "indexed",
                        "documentId": str(document_id),
                        "chunks": len(chunks),
                    }
                )
            except Exception as error:
                await self.db.rollback()
                summary["failedFiles"] += 1
                summary["files"].append(
                    {"sourcePath": file.source_path, "status": "failed", "reason": str(error)}
                )
        return summary

    def _build_chunks(self, file: RawDocument, tenant: str) -> list[dict[str, Any]]:
        chunks = []
        for segment in self.loader.load(file):
            for split in self.splitter.split_text(segment["text"]):
                content = split.strip()
                if content:
                    chunks.append(
                        {
                            "content": content,
                            "metadata": {
                                **segment["metadata"],
                                "sourcePath": file.source_path,
                                "fileName": file.file_name,
                                "fileType": file.file_type,
                                "type": file.document_type,
                                "tenantCode": tenant,
                            },
                        }
                    )
        return chunks

    async def _embed(self, texts: list[str]) -> list[list[float]]:
        result: list[list[float]] = []
        size = min(max(self.settings.rag_embedding_batch_size, 1), 10)
        for index in range(0, len(texts), size):
            result.extend(await self.embeddings.aembed_documents(texts[index : index + size]))
        return result

    async def _replace(
        self,
        file: RawDocument,
        tenant: str,
        chunks: list[dict[str, Any]],
        vectors: list[list[float]],
    ) -> int:
        result = await self.db.execute(
            text("""INSERT INTO aplus_documents (tenant_code, source_path, file_name, file_type, "type", tag, file_hash, file_size, status, updated_at)
                    VALUES (:tenant, :source, :name, :file_type, :document_type, '', :file_hash, :file_size, 'processing', CURRENT_TIMESTAMP)
                    ON CONFLICT (tenant_code, source_path) DO UPDATE SET file_name=EXCLUDED.file_name, file_type=EXCLUDED.file_type, "type"=EXCLUDED."type", file_hash=EXCLUDED.file_hash, file_size=EXCLUDED.file_size, status='processing', error=NULL, updated_at=CURRENT_TIMESTAMP
                    RETURNING id"""),
            {
                "tenant": tenant,
                "source": file.source_path,
                "name": file.file_name,
                "file_type": file.file_type,
                "document_type": file.document_type,
                "file_hash": file.file_hash,
                "file_size": file.file_size,
            },
        )
        document_id = int(result.scalar_one())
        await self.db.execute(
            text("DELETE FROM aplus_chunks WHERE tenant_code=:tenant AND document_id=:document_id"),
            {"tenant": tenant, "document_id": document_id},
        )
        for index, chunk in enumerate(chunks):
            await self.db.execute(
                text(
                    "INSERT INTO aplus_chunks (tenant_code, document_id, chunk_index, content, metadata, embedding) VALUES (:tenant, :document_id, :chunk_index, :content, CAST(:metadata AS jsonb), CAST(:embedding AS vector))"
                ),
                {
                    "tenant": tenant,
                    "document_id": document_id,
                    "chunk_index": index,
                    "content": chunk["content"],
                    "metadata": json.dumps(chunk["metadata"], ensure_ascii=False),
                    "embedding": json.dumps(vectors[index]),
                },
            )
        await self.db.execute(
            text(
                "UPDATE aplus_documents SET chunk_count=:count, status='ready', error=NULL, updated_at=CURRENT_TIMESTAMP, ingested_at=CURRENT_TIMESTAMP WHERE id=:document_id AND tenant_code=:tenant"
            ),
            {"count": len(chunks), "document_id": document_id, "tenant": tenant},
        )
        await self.db.commit()
        return document_id

    @staticmethod
    def _unique_target(root: Path, file_name: str) -> Path:
        stem, suffix = Path(file_name).stem, Path(file_name).suffix
        for index in range(10_000):
            target = root / (file_name if index == 0 else f"{stem} ({index}){suffix}")
            if not target.exists():
                return target
        raise HTTPException(400, "too many files with the same name")
