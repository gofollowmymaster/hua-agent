from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from fastapi import HTTPException

from python_api.core.config import Settings, get_settings

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".ppt", ".pptx", ".txt", ".md"}

@dataclass(slots=True)
class RawDocument:
    """Describes a source file before text extraction."""

    absolute_path: Path
    source_path: str
    file_name: str
    file_type: str
    document_type: str
    file_size: int
    file_hash: str


class DocumentLoader:
    """Extracts supported office and text formats into metadata-bearing segments."""

    def __init__(self, settings: Settings | None = None) -> None:
        self.settings = settings or get_settings()

    def raw_root(self, tenant: str) -> Path:
        """Returns the configured tenant raw directory."""
        return (self.settings.workspace_root / self.settings.raw_dir / tenant).resolve()

    def scan(self, tenant: str) -> list[RawDocument]:
        """Recursively describes supported tenant files."""
        root = self.raw_root(tenant)
        if not root.exists():
            return []
        return [
            self.describe(path, root, tenant)
            for path in root.rglob("*")
            if path.is_file()
            and path.name not in {".DS_Store", ".WeDrive"}
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
        ]

    def describe(
        self, path: Path, root: Path, tenant: str, file_name: str | None = None
    ) -> RawDocument:
        """Collects stable path, type, size and SHA-256 metadata."""
        name = file_name or path.name
        extension = Path(name).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise HTTPException(400, f"Unsupported file type: {extension or 'none'}")
        source_path = path.relative_to(root).as_posix()
        parts = Path(source_path).parts
        return RawDocument(
            absolute_path=path,
            source_path=source_path,
            file_name=name,
            file_type=extension.removeprefix("."),
            document_type=parts[0] if len(parts) > 1 else "",
            file_size=path.stat().st_size,
            file_hash=hashlib.sha256(path.read_bytes()).hexdigest(),
        )

    def load(self, document: RawDocument) -> list[dict[str, Any]]:
        """Loads one supported file and returns text segments."""
        suffix = document.absolute_path.suffix.lower()
        if suffix in {".txt", ".md"}:
            return [{"text": document.absolute_path.read_text(encoding="utf-8"), "metadata": {}}]
        if suffix == ".pdf":
            from pypdf import PdfReader

            return [
                {
                    "text": self._normalize_pdf(page.extract_text() or ""),
                    "metadata": {"page": index + 1},
                }
                for index, page in enumerate(PdfReader(document.absolute_path).pages)
            ]
        if suffix == ".docx":
            from docx import Document

            loaded = Document(str(document.absolute_path))
            text_content = "\n".join(paragraph.text for paragraph in loaded.paragraphs)
            return [{"text": text_content, "metadata": {}}]
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(document.absolute_path, data_only=True, read_only=True)
            segments = []
            for sheet in workbook.worksheets:
                rows = []
                for row_index, values in enumerate(sheet.iter_rows(values_only=True), 1):
                    cells = [
                        str(value).strip()
                        for value in values
                        if value is not None and str(value).strip()
                    ]
                    if cells:
                        rows.append(f"第 {row_index} 行：" + " | ".join(cells))
                segments.append({"text": "\n".join(rows), "metadata": {"sheet": sheet.title}})
            return segments
        if suffix == ".pptx":
            from pptx import Presentation

            presentation = Presentation(str(document.absolute_path))
            return [
                {
                    "text": "\n".join(
                        shape.text for shape in slide.shapes if hasattr(shape, "text")
                    ),
                    "metadata": {"page": index + 1},
                }
                for index, slide in enumerate(presentation.slides)
            ]
        if suffix == ".ppt":
            content = document.absolute_path.read_bytes().decode("latin1", errors="ignore")
            extracted = re.sub(r"[^\x20-\x7e\u4e00-\u9fa5]+", " ", content)
            extracted = re.sub(r"\s+", " ", extracted).strip()
            if not extracted:
                raise HTTPException(
                    400,
                    "Legacy .ppt file has no extractable text. Convert it to .pptx or PDF and retry.",
                )
            return [{"text": extracted, "metadata": {"legacyPpt": True}}]
        raise HTTPException(400, f"Unsupported file type: {suffix}")

    @staticmethod
    def _normalize_pdf(value: str) -> str:
        """Normalizes extracted PDF whitespace while preserving paragraph breaks."""
        value = value.replace("\r\n", "\n")
        value = re.sub(
            r"\b(?:[A-Za-z][ \t]+)+[A-Za-z]\b",
            lambda match: re.sub(r"[ \t]+", "", match.group()),
            value,
        )
        value = re.sub(r"[ \t]+", " ", value)
        return re.sub(r"\n{3,}", "\n\n", value).strip()
